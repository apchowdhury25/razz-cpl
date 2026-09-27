from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any, Sequence

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.money import money

BDT_EXCEL_FORMAT = r"[>=10000000]##\,##\,##\,##0.00;[>=100000]##\,##\,##0.00;##,##0.00"


def excel_response(
    *,
    title: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    filename: str,
    company: str = "Razz CNPL",
    filters: str = "",
) -> StreamingResponse:
    wb = Workbook()
    ws = wb.active
    safe_title = "".join("-" if ch in r"[]:*?/\\" else ch for ch in title).strip() or "Report"
    ws.title = safe_title[:31]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0F3D3E")
    title_font = Font(bold=True, size=14, color="0F3D3E")
    money_font = Font(name="Calibri")
    thin = Border(
        left=Side(style="thin", color="D9D3C7"),
        right=Side(style="thin", color="D9D3C7"),
        top=Side(style="thin", color="D9D3C7"),
        bottom=Side(style="thin", color="D9D3C7"),
    )

    ws["A1"] = company
    ws["A1"].font = title_font
    ws["A2"] = title
    ws["A2"].font = Font(bold=True, size=12)
    ws["A3"] = filters
    ws["A4"] = f"Generated: {datetime.now().strftime('%d-%m-%Y %H:%M')}"

    start_row = 6
    for col, header in enumerate(headers, 1):
        cell = ws.cell(start_row, col, header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin

    for r_idx, row in enumerate(rows, start_row + 1):
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(r_idx, c_idx, value)
            cell.border = thin
            cell.font = money_font
            if isinstance(value, (int, float)) or hasattr(value, "as_tuple"):
                try:
                    quantized = money(value)
                    cell.value = float(quantized)
                    cell.number_format = BDT_EXCEL_FORMAT
                    cell.alignment = Alignment(horizontal="right")
                except Exception:
                    cell.value = str(value)

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = max(14, min(36, len(headers[col - 1]) + 8))

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:6"

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    if not filename.endswith(".xlsx"):
        filename += ".xlsx"
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
